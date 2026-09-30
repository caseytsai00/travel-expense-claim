"""Parse the five retained Project D sources into normalized in-memory records.

Every function here reads ONLY from files listed in a run's sources.json — no
case-specific values (claim IDs, amounts, employee names) are hard-coded.
Only the policy's own stated business constants (director threshold, filing
window) live in policy_rules.py, since those are policy text, not case data.
"""
import csv
import json
import re
from pathlib import Path

import openpyxl
import PyPDF2


def load_sources_manifest(sources_json_path):
    return json.loads(Path(sources_json_path).read_text())


def _resolve(run_root, path):
    """sources.json paths are repo-root-relative; resolve against run_root's parent chain."""
    p = Path(path)
    if p.is_absolute():
        return p
    # sources.json 'path' fields are relative to the repo root, and run_root is
    # .../artifacts/runs/<run-id>, i.e. repo_root/artifacts/runs/<run-id>
    repo_root = Path(run_root).parent.parent.parent
    return repo_root / path


# ---------------------------------------------------------------------------
# Source 1: policy (Notion, retained as markdown) — not machine-parsed here.
# Its numeric/business rules are encoded by hand in policy_rules.py, with each
# constant citing the exact sentence in artifacts/runs/<run>/sources/policy.md
# it came from. This function just exposes the retained text + its source_id
# for citation purposes.
# ---------------------------------------------------------------------------

def load_policy_text(manifest, run_root):
    entry = next(f for f in manifest["files"] if f["source_id"] == "SRC-1-policy")
    return entry["source_id"], _resolve(run_root, entry["path"]).read_text()


# ---------------------------------------------------------------------------
# Source 2 + 3: claim binders (PDF, one claim revision per page)
# ---------------------------------------------------------------------------

_CLAIM_FIELD_RE = {
    "claim_id": r"Claim reference:\s*(\S+)",
    "revision": r"Revision:\s*(\d+)",
    "arrival_batch": r"Arrival batch:\s*(\d+)",
    "employee": r"Employee:\s*(\S+)",
    "trip_id": r"Trip reference:\s*(\S+)",
    "trip_revision": r"Trip revision:\s*(\d+)",
    "submitted_date": r"Submitted date:\s*(\S+)",
    "submission_status": r"Status:\s*(\S+)",
}

_LINE_RE = re.compile(
    r"\n(\d+)\n(COST-\S+)\n(R-COST-\S+)\n(\d{4}-\d{2}-\d{2})\n([\d.]+)\n([A-Z]{3})"
)

_EXCEPTION_RE = re.compile(
    r"Finance exception\n"
    r"Authorizing Finance officer:\s*(\S+)\n"
    r"Applies to claim:\s*(\S+)\n"
    r"Applies to revision:\s*(\d+)\n"
    r"Expense reference:\s*(\S+)\n"
    r"Allowed original amount:\s*([\d.]+)\n"
    r"Covered issues:\s*(.+?)\n"
    r"(?:Travel cancellation reference:\s*(\S+)\n)?"
    r"Exception reason:\s*(.+?)(?:\n\n|\Z)",
    re.S,
)


def _parse_claim_page(text, source_id, page_no):
    d = {"source_id": source_id, "page": page_no, "raw_text": text}
    for key, pat in _CLAIM_FIELD_RE.items():
        m = re.search(pat, text)
        d[key] = m.group(1) if m else None
    for intkey in ("revision", "arrival_batch", "trip_revision"):
        if d.get(intkey) is not None:
            d[intkey] = int(d[intkey])

    reason_m = re.search(r"Claim reason:\s*(.+?)\nRelated earlier claims:", text, re.S)
    d["claim_reason"] = reason_m.group(1).strip() if reason_m else None
    rel_m = re.search(r"Related earlier claims:\s*(.+?)\nExpense details", text, re.S)
    rel = rel_m.group(1).strip() if rel_m else "None"
    d["related_earlier_claims"] = [] if rel == "None" else [x.strip() for x in rel.split(",")]

    lines = []
    for m in _LINE_RE.finditer(text):
        lines.append(
            {
                "line": int(m.group(1)),
                "cost_id": m.group(2),
                "receipt_id": m.group(3),
                "first_submitted_date": m.group(4),
                "original_amount": float(m.group(5)),
                "currency": m.group(6),
            }
        )
    d["lines"] = lines

    exceptions = []
    for m in _EXCEPTION_RE.finditer(text):
        covered = m.group(6).strip()
        exceptions.append(
            {
                "actor": m.group(1),
                "claim_id": m.group(2),
                "revision": int(m.group(3)),
                "cost_id": m.group(4),
                "allowed_original_amount": float(m.group(5)),
                "covered_issues": [] if covered.lower() == "none" else [
                    c.strip().lower().replace(" ", "_") for c in covered.split(",")
                ],
                "travel_cancellation_ref": m.group(7),
                "exception_reason": m.group(8).strip(),
            }
        )
    d["exceptions"] = exceptions
    return d


def load_claim_pages(manifest, run_root):
    """Returns a flat list of claim-revision records, one per PDF page, across
    both the initial binder and the updates binder, each tagged with its own
    source_id and page number for citation."""
    pages = []
    for source_id in ("SRC-2-claims-initial", "SRC-3-claims-updates"):
        entry = next(f for f in manifest["files"] if f["source_id"] == source_id)
        path = _resolve(run_root, entry["path"])
        reader = PyPDF2.PdfReader(str(path))
        for i, page in enumerate(reader.pages):
            pages.append(_parse_claim_page(page.extract_text(), source_id, i + 1))
    return pages


def group_claims_by_id(claim_pages):
    """Groups flat page records into {claim_id: [revision_record, ...]} sorted
    by revision number, so callers can walk each claim's revision history."""
    by_id = {}
    for rec in claim_pages:
        by_id.setdefault(rec["claim_id"], []).append(rec)
    for cid in by_id:
        by_id[cid].sort(key=lambda r: r["revision"])
    return by_id


# ---------------------------------------------------------------------------
# Source 4: receipts (PDF, one receipt per page)
# ---------------------------------------------------------------------------

_RECEIPT_FIELD_RE = {
    "receipt_id": r"Receipt reference:\s*(\S+)",
    "cost_id": r"Expense reference:\s*(\S+)",
    "employee": r"Employee:\s*(\S+)",
    "trip_id": r"Trip reference:\s*(\S+)",
    "category": r"Category:\s*(\S+)",
    "currency": r"Currency:\s*(\S+)",
    "issued_date": r"Issued date:\s*(\S+)",
}


def load_receipts(manifest, run_root):
    entry = next(f for f in manifest["files"] if f["source_id"] == "SRC-4-receipts")
    path = _resolve(run_root, entry["path"])
    reader = PyPDF2.PdfReader(str(path))
    receipts = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        d = {"source_id": "SRC-4-receipts", "page": i + 1}
        for key, pat in _RECEIPT_FIELD_RE.items():
            m = re.search(pat, text)
            d[key] = m.group(1) if m else None
        gm = re.search(r"Gross amount:\s*([\d.]+)", text)
        d["gross_amount"] = float(gm.group(1)) if gm else None
        um = re.search(r"Documented units:\s*(\S+)", text)
        d["documented_units"] = int(um.group(1)) if um and um.group(1).isdigit() else None
        desc = re.search(r"Description:\s*(.+?)\nIssued date:", text, re.S)
        d["description"] = desc.group(1).strip() if desc else None
        receipts.append(d)
    return receipts


# ---------------------------------------------------------------------------
# Source 5: travel-registers workbook (xlsx, 10 tabs)
# ---------------------------------------------------------------------------

def _sheet_to_dicts(ws):
    """Rows 1-4 of every tab are title/note/blank/blank; row 5 is the header."""
    rows = list(ws.iter_rows(values_only=True))
    header = rows[4]
    out = []
    for row in rows[5:]:
        if all(v is None for v in row):
            continue
        out.append(dict(zip(header, row)))
    return out


def load_workbook(manifest, run_root):
    entry = next(f for f in manifest["files"] if f["source_id"] == "SRC-5-travel-registers")
    path = _resolve(run_root, entry["path"])
    wb = openpyxl.load_workbook(str(path), data_only=True)
    tabs = {}
    for ws in wb.worksheets:
        tabs[ws.title] = _sheet_to_dicts(ws)
    return tabs


def load_all(sources_json_path):
    manifest = load_sources_manifest(sources_json_path)
    run_root = Path(sources_json_path).parent
    policy_source_id, policy_text = load_policy_text(manifest, run_root)
    claim_pages = load_claim_pages(manifest, run_root)
    return {
        "manifest": manifest,
        "policy_source_id": policy_source_id,
        "policy_text": policy_text,
        "claims_by_id": group_claims_by_id(claim_pages),
        "receipts": load_receipts(manifest, run_root),
        "tabs": load_workbook(manifest, run_root),
    }
